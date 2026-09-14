"""
Workplace-Pulse-Telemetry: Enterprise Sample Corpus Generator
Engineered by Fabio Torres (neurodeveloper11)

Generates realistic mock enterprise files (DOCX, PDF, XLSX, EML, TXT)
organized in realistic departmental hierarchies with sensitive PII and psychometric markers.
"""

import os
import zipfile
from pathlib import Path
from datetime import datetime

import docx
from docx.shared import Inches, Pt, RGBColor
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors


def create_sample_docx(target_path: Path):
    """Creates a sample meeting minutes Word document (.docx)."""
    target_path.parent.mkdir(parents=True, exist_ok=True)
    doc = docx.Document()

    # Document properties
    doc.core_properties.title = "Acta de Comité de Convivencia Laboral - Sesión Ordinaria N° 03"
    doc.core_properties.author = "Lic. Andrea Morales (Secretaría Técnica)"
    doc.core_properties.created = datetime(2026, 3, 10, 14, 30)

    title = doc.add_heading("ACTA DE COMITÉ DE CONVIVENCIA LABORAL", level=1)
    doc.add_paragraph("Fecha: 10 de Marzo de 2026 | Lugar: Sala de Juntas Operaciones Portuarias")
    doc.add_paragraph("Empresa: Terminal Portuaria del Pacífico S.A.S.")

    doc.add_heading("1. ASISTENCIA Y VERIFICACIÓN DEL QUÓRUM", level=2)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Nombre"
    hdr_cells[1].text = "Cargo"
    hdr_cells[2].text = "Contacto / Cédula"

    attendees = [
        ("Dr. Fernando Rivas", "Presidente del Comité", "CC 14.890.123 - fernando.rivas@puerto.com"),
        ("Carlos Pérez", "Representante de Operarios", "CC 11.234.567 - Cel: +57 315 456 7890"),
        ("Ing. Sofía Gómez", "Supervisora de Operaciones", "CC 31.456.789 - sofia.gomez@puerto.com"),
        ("Lic. Andrea Morales", "Secretaria Técnica SST", "CC 29.876.543 - andrea.morales@puerto.com")
    ]
    for name, role, contact in attendees:
        row_cells = table.add_row().cells
        row_cells[0].text = name
        row_cells[1].text = role
        row_cells[2].text = contact

    doc.add_heading("2. ORDEN DEL DÍA", level=2)
    doc.add_paragraph("1. Revisión de quejas por sobrecarga y turnos dobles en muelle.")
    doc.add_paragraph("2. Casos de fricción comunicacional entre jefatura de turno y operadores de grúa.")
    doc.add_paragraph("3. Acuerdos y compromisos para cumplimiento de la Resolución 2764/2022.")

    doc.add_heading("3. DESARROLLO DE LA REUNIÓN E INTERVENCIONES", level=2)
    interventions = [
        "Presidente: Damos inicio a la sesión. Escuchamos primero la exposición de los representantes de patio y muelle.",
        "Carlos Pérez (Operaciones): Me encuentro totalmente agotado y al límite con el insomnio. En las últimas tres semanas hemos trabajado jornadas continuas de 16 horas sin compensatorio. La sobrecarga física y mental es insostenible.",
        "Ing. Sofía Gómez: Como ya te había dicho en la reunión anterior, la programación de los buques portacontenedores es obligatoria y no admite demoras. No es mi trabajo resolver la falta de personal en los muelles.",
        "Carlos Pérez (Operaciones): Otra vez con lo mismo. No vuelvo a repetir que la seguridad en las grúas pórtico no se puede descuidar por apagar incendios.",
        "Ing. Sofía Gómez: Pido disculpas si mi tono sonó agresivo. Cometí un error al asumir que contaban con apoyo del turno nocturno. ¿Qué opinan si redistribuimos las cuadrillas de guardia?",
        "Presidente: Excelente propuesta. Agradezco la disposición al diálogo constructivo. Propongo fijar una rotación formal con descanso mínimo garantizado de 12 horas."
    ]
    for item in interventions:
        doc.add_paragraph(item)

    doc.add_heading("4. CONCLUSIONES Y COMPROMISOS", level=2)
    doc.add_paragraph("Compromiso 1. Talento Humano revisará el pago de horas extras adeudadas (salario base de $4.800.000 COP más recargos) antes del 20 de marzo.")
    doc.add_paragraph("Compromiso 2. Se prohíben las comunicaciones operativas por WhatsApp o correo entre las 19:00 y las 07:00, amparados en la Ley de Desconexión Laboral.")

    doc.save(str(target_path))


def create_sample_pdf(target_path: Path):
    """Creates a sample incident and audit report PDF."""
    target_path.parent.mkdir(parents=True, exist_ok=True)
    pdf_doc = SimpleDocTemplate(str(target_path), pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    title_style = ParagraphStyle(
        'TitleStyle',
        parent=styles['Heading1'],
        fontSize=16,
        textColor=colors.HexColor('#0f172a'),
        spaceAfter=12
    )
    body_style = ParagraphStyle(
        'BodyStyle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        spaceAfter=8
    )
    alert_style = ParagraphStyle(
        'AlertStyle',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#991b1b'),
        spaceAfter=8
    )

    story.append(Paragraph("INFORME PERICIAL DE AUDITORÍA Y CLIMA OCUPACIONAL", title_style))
    story.append(Paragraph("<b>Área Auditada:</b> Finanzas y Contabilidad | <b>Fecha:</b> 12 de Febrero de 2026", body_style))
    story.append(Paragraph("<b>Auditor Responsable:</b> Dr. Fabio Torres (Licencia SST) | IP Servidor: 10.14.0.22", body_style))
    story.append(Spacer(1, 10))

    story.append(Paragraph("<b>1. RESUMEN EJECUTIVO DE HALLAZGOS PSICOSOCIALES</b>", styles['Heading2']))
    story.append(Paragraph(
        "Se evaluaron 28 colaboradores del departamento financiero. Se evidencia alta demanda cuantitativa asociada al cierre fiscal anual. "
        "Varios analistas expresan: <i>'Llevo dos semanas sin dormir, estoy drenada y con saturación extrema por los reportes inmediatos'</i>.",
        body_style
    ))

    story.append(Paragraph(
        "Adicionalmente, se identificaron incidentes de comunicación hostil dirigidos a la analista María Rodríguez (CC 52.123.456, salario $5.200.000 COP, email maria.rodriguez@finanzas.com). "
        "El supervisor inmediato envió directivas con tono imperativo: <i>'Es obligatorio enviar ya los balances, sin excepciones ni quejas'</i>.",
        alert_style
    ))

    story.append(Paragraph("<b>2. MATRIZ DE RIESGO DE DEMANDA Y CONTROL (KARASEK)</b>", styles['Heading2']))
    story.append(Paragraph(
        "El equipo clasifica en el cuadrante de <b>Alta Tensión (High Strain)</b>: alta exigencia en plazos de entrega y nulo margen de autonomía para reorganizar prioridades. "
        "Recomendación: Establecer mesas de nivelación y suprimir llamadas fuera del horario contractual (desconexión digital).",
        body_style
    ))

    pdf_doc.build(story)


def create_sample_excel(target_path: Path):
    """Creates a sample shift and incident Excel log (.xlsx)."""
    target_path.parent.mkdir(parents=True, exist_ok=True)

    data = {
        "Fecha": ["2026-03-01 22:15", "2026-03-02 03:30", "2026-03-03 11:00", "2026-03-04 18:45", "2026-03-05 23:50"],
        "Colaborador": ["Dev_Juan_Camilo", "Dev_Laura_Vargas", "Lead_Marcos_Paz", "Dev_Juan_Camilo", "Dev_Laura_Vargas"],
        "Incidente_ID": ["INC-901", "INC-902", "INC-903", "INC-904", "INC-905"],
        "Descripcion_Observaciones": [
            "URGENTE: Se cayó la base de datos de producción en la IP 192.168.1.105. Llevo 14 horas de guardia continua y estoy quemado, necesito relevo ya!",
            "Cometí un error en el script de migración, pido disculpas. Ya lo revertí y el cluster está estable. ¿Alguien tiene dudas o sugerencias para el post-mortem?",
            "Revisión de sprint: Gran trabajo equipo con el lanzamiento de la nueva API, felicitaciones a todos por el esfuerzo coordinado.",
            "Como ya te había dicho en Slack, el despliegue del módulo de facturación no es mi trabajo. Favor leer el manual antes de asignarme el ticket.",
            "Alerta de desborde de memoria a medianoche. Apagando incendios otra vez en deshoras. Estoy agotada y sin energía para continuar el turno de mañana."
        ],
        "Horas_Extra_Reportadas": [4, 5, 0, 2, 6]
    }

    df = pd.DataFrame(data)
    with pd.ExcelWriter(target_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Guardias_Incidentes", index=False)


def create_sample_email(target_path: Path):
    """Creates an RFC-822 formatted email file (.eml)."""
    target_path.parent.mkdir(parents=True, exist_ok=True)

    eml_content = """From: "Dra. Patricia Salazar" <patricia.salazar@terminalportuaria.com>
To: "Comite SST" <sst_comite@terminalportuaria.com>
Subject: Alerta Psicosocial: Aumento de friccion y agotamiento en turno nocturno
Date: Mon, 09 Mar 2026 21:45:00 -0500
Content-Type: text/plain; charset="utf-8"

Estimados miembros del Comité de Convivencia y SST:

Les escribo con carácter de urgencia. Durante la inspección del turno de noche en muelles, varios colaboradores con cédulas CC 10.987.654 y CC 11.876.543 manifestaron encontrarse en estado de agotamiento severo e insomnio acumulado.

Uno de los operadores expresó textualmente: "No doy más, estamos al límite y no es justo que nos llamen a deshoras los fines de semana". 
Además, se percibió hostilidad entre supervisores y operarios con frases como "no me pagan lo suficiente para aguantar esto" y "hagan lo que quieran".

Solicito convocar a una sesión extraordinaria esta semana para revisar las directrices de la Resolución 2764/2022 y aplicar los protocolos de seguridad psicológica de Amy Edmondson.

Atentamente,
Dra. Patricia Salazar
Directora de Talento Humano y Bienestar Laboral
Teléfono: +57 320 654 9870
Terminal Portuaria del Pacífico
"""
    with open(target_path, "w", encoding="utf-8") as f:
        f.write(eml_content)


def create_sample_txt(target_path: Path):
    """Creates a sample plain text complaint / memo."""
    target_path.parent.mkdir(parents=True, exist_ok=True)

    content = """ACTA DE DESCARGO LABORAL Y SOLICITUD DE REASIGNACIÓN
Fecha: 05 de Marzo de 2026
Departamento: Atencion_Cliente
Colaborador: Diana Marcela Restrepo (Cédula: CC 38.456.123)

HECHOS EXPUESTOS:
Por medio de la presente, solicito formalmente la intervención del área de Talento Humano.
Durante el último mes, el volumen de llamadas y reclamos ha sobrepasado los límites contractuales. 
Tengo que atender más de 90 clientes diarios sin pausa activa ni tiempo para almuerzo.
El supervisor nos exige: "Debes cumplir la métrica sin excepciones, está prohibido levantarse del puesto".

SITUACIÓN EMOCIONAL:
Me siento totalmente desbordada, con crisis de ansiedad y migrañas continuas. 
No me siento con la seguridad psicológica para expresar mis inquietudes porque temo represalias laborales.
Agradezco que se realice una evaluación con el Psicólogo de la empresa bajo el protocolo de la Resolución 2764/2022.

Firma del Trabajador: Diana M. Restrepo
Contacto: Celular 318 432 1098 | diana.restrepo@empresa.com
Salario Devengado: $2.800.000 COP
"""
    with open(target_path, "w", encoding="utf-8") as f:
        f.write(content)


def generate_all_samples(base_dir: Path) -> Path:
    """Generates all individual files and packages them into a clean ZIP archive."""
    corpus_dir = base_dir / "sample_documents"
    corpus_dir.mkdir(parents=True, exist_ok=True)

    docx_file = corpus_dir / "Operaciones_Portuarias" / "2026" / "Actas_Comite" / "acta_convivencia_marzo_2026.docx"
    pdf_file = corpus_dir / "Finanzas_Contabilidad" / "2026" / "Auditorias" / "informe_auditoria_financiera_2026.pdf"
    xlsx_file = corpus_dir / "Ingenieria_Core" / "2026" / "Bitacoras" / "bitacora_guardias_incidentes.xlsx"
    eml_file = corpus_dir / "Comite_Convivencia_SST" / "2026" / "Emails" / "alerta_friccion_turno_nocturno.eml"
    txt_file = corpus_dir / "Atencion_Cliente" / "2026" / "Quejas" / "descargo_solicitud_apoyo.txt"

    create_sample_docx(docx_file)
    create_sample_pdf(pdf_file)
    create_sample_excel(xlsx_file)
    create_sample_email(eml_file)
    create_sample_txt(txt_file)

    # Package into a ZIP archive for demonstration in UI and testing
    zip_path = corpus_dir / "auditoria_organizacional_2026_demo.zip"
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for file_path in [docx_file, pdf_file, xlsx_file, eml_file, txt_file]:
            arcname = file_path.relative_to(corpus_dir)
            zf.write(file_path, arcname=str(arcname))

    print(f"Sample corpus generated successfully in: {corpus_dir}")
    print(f"Sample ZIP archive created at: {zip_path}")
    return zip_path


if __name__ == "__main__":
    current_dir = Path(__file__).resolve().parent
    generate_all_samples(current_dir)
